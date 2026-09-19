import os
from datetime import date

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from .config import settings
from .models import Finding, SlidePlan

BRAND_PRIMARY = RGBColor(0x1F, 0x3A, 0x5F)
BRAND_ACCENT = RGBColor(0x00, 0xB8, 0x9E)
BRAND_TEXT = RGBColor(0x22, 0x22, 0x22)

LOGO_PATH = os.path.join(os.path.dirname(__file__), "..", "assets", "branding", "logo.png")


def _style_title(shape, size: int = 32) -> None:
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = True
            run.font.color.rgb = BRAND_PRIMARY


def _add_footer(slide, text: str, prs: Presentation) -> None:
    if not text:
        return
    left = Inches(0.5)
    top = prs.slide_height - Inches(0.5)
    width = prs.slide_width - Inches(1)
    box = slide.shapes.add_textbox(left, top, width, Inches(0.4))
    frame = box.text_frame
    frame.text = text
    paragraph = frame.paragraphs[0]
    paragraph.font.size = Pt(10)
    paragraph.font.color.rgb = BRAND_ACCENT
    paragraph.alignment = PP_ALIGN.LEFT


def build_deck(
    slide_plan: SlidePlan,
    findings: list[Finding],
    output_path: str,
) -> None:
    prs = Presentation()

    ref_numbers: dict[str, int] = {}
    findings_by_id = {f.paper_id: f for f in findings}

    def ref_number_for(paper_id: str) -> int:
        if paper_id not in ref_numbers:
            ref_numbers[paper_id] = len(ref_numbers) + 1
        return ref_numbers[paper_id]

    # Title slide
    title_layout = prs.slide_layouts[0]
    title_slide = prs.slides.add_slide(title_layout)
    title_slide.shapes.title.text = slide_plan.deck_title
    _style_title(title_slide.shapes.title, size=40)
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = f"Research synthesis · Generated {date.today().isoformat()}"
    if os.path.exists(LOGO_PATH):
        title_slide.shapes.add_picture(LOGO_PATH, prs.slide_width - Inches(1.5), Inches(0.3), height=Inches(1))

    # Content slides
    content_layout = prs.slide_layouts[1]
    for slide_data in slide_plan.slides:
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = slide_data.title
        _style_title(slide.shapes.title, size=28)

        body = slide.placeholders[1]
        text_frame = body.text_frame
        text_frame.clear()
        for index, bullet in enumerate(slide_data.bullets):
            paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
            paragraph.text = bullet
            paragraph.font.size = Pt(20)
            paragraph.font.color.rgb = BRAND_TEXT
            paragraph.level = 0

        if slide_data.citations:
            numbers = sorted({ref_number_for(pid) for pid in slide_data.citations if pid in findings_by_id})
            footer_text = "Sources: " + " ".join(f"[{n}]" for n in numbers)
            _add_footer(slide, footer_text, prs)

        notes_slide = slide.notes_slide
        notes_slide.notes_text_frame.text = slide_data.speaker_notes

    # References slide
    if ref_numbers:
        ref_layout = prs.slide_layouts[1]
        ref_slide = prs.slides.add_slide(ref_layout)
        ref_slide.shapes.title.text = "References"
        _style_title(ref_slide.shapes.title, size=28)

        body = ref_slide.placeholders[1]
        text_frame = body.text_frame
        text_frame.clear()

        ordered = sorted(ref_numbers.items(), key=lambda item: item[1])
        for index, (paper_id, number) in enumerate(ordered):
            finding = findings_by_id.get(paper_id)
            if finding is None:
                continue
            citation = f"[{number}] {finding.title} — {finding.authors} ({finding.year or 'n.d.'})"
            paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
            paragraph.text = citation
            paragraph.font.size = Pt(14)
            paragraph.font.color.rgb = BRAND_TEXT

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    prs.save(output_path)
